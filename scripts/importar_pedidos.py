"""Importa la hoja STATUS desde un .xlsx descargado a la tabla `pedidos`.

Respaldo para cuando no hay acceso al Google Sheet por API. Lo normal es
usar scripts/sincronizar_pedidos.py, que lee la hoja en vivo.

Reemplaza la tabla completa: la hoja es la fuente de verdad y esta copia
solo sirve para que el bot conteste al instante.

Uso:
    uv run python scripts/importar_pedidos.py "PACIENTES SUPERVISION.xlsx"
    uv run python scripts/importar_pedidos.py archivo.xlsx --meses 6
    uv run python scripts/importar_pedidos.py archivo.xlsx --hoja STATUS
"""

import argparse
import asyncio
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent / "src"))

from sidhe_agent.config import get_settings
from sidhe_agent.db.session import dispose_engine
from sidhe_agent.services import pedidos


def leer_hoja(ruta: Path, hoja: str) -> list[tuple]:
    try:
        import openpyxl
    except ImportError:
        raise SystemExit(
            "Falta openpyxl. Córrelo así: "
            "uv run --with openpyxl python scripts/importar_pedidos.py ..."
        ) from None

    libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    if hoja not in libro.sheetnames:
        raise SystemExit(
            f"La hoja '{hoja}' no existe. Hojas disponibles: {libro.sheetnames[:12]}"
        )
    filas = list(libro[hoja].iter_rows(values_only=True))
    libro.close()
    if not filas:
        raise SystemExit("La hoja está vacía")

    faltantes = pedidos.columnas_faltantes(list(filas[0]))
    if faltantes:
        raise SystemExit(f"El encabezado no tiene {faltantes}: {filas[0][:9]}")
    return filas[1:]


def imprimir_resumen(registros: list[dict], en_hoja: int, meses: int) -> None:
    datos = pedidos.resumen(registros)
    total = datos["pedidos"]
    print(f"Filas en la hoja: {en_hoja}")
    print(f"Pedidos importados (últimos {meses} meses): {total}")
    print(f"Rango de fechas: {datos['desde']} a {datos['hasta']}")
    print(
        f"Con teléfono (búsqueda automática): {datos['con_telefono']} "
        f"({100 * datos['con_telefono'] // total}%)"
    )
    print("Por categoría:")
    for categoria, n in sorted(
        datos["por_categoria"].items(), key=lambda x: -x[1]
    ):
        print(f"  {n:7}  {categoria}")

    revision = [r for r in registros if r["status"] == pedidos.REVISION]
    if revision:
        print("\nStatus que el bot mandará a un asesor (los más comunes):")
        for texto, n in Counter(
            r["status_original"] for r in revision
        ).most_common(5):
            print(f"  {n:7}  {texto!r}")


async def main(ruta: str, hoja: str, meses: int) -> None:
    archivo = Path(ruta)
    if not archivo.is_file():
        raise SystemExit(f"No existe el archivo: {archivo}")

    crudas = leer_hoja(archivo, hoja)
    registros = pedidos.preparar_filas(crudas, meses=meses)
    if not registros:
        raise SystemExit(
            f"Ninguna fila con nombre quedó dentro de los últimos {meses} meses"
        )

    await pedidos.reemplazar(registros)
    imprimir_resumen(registros, len(crudas), meses)
    await dispose_engine()


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("archivo", help="ruta del .xlsx de operaciones")
    parser.add_argument("--hoja", default="STATUS")
    parser.add_argument(
        "--meses",
        type=int,
        default=get_settings().pedidos_meses_historial,
        help="meses de historial a conservar (0 = todo)",
    )
    args = parser.parse_args()
    asyncio.run(main(args.archivo, args.hoja, args.meses))
